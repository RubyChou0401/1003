"""Excel 匯入：解析 → 檢核 → 預覽 → 確認後才寫入。錯誤列永遠不會寫入正式資料。"""
import io
import json
import re
from datetime import date, datetime, timedelta

import openpyxl
from werkzeug.security import generate_password_hash

import db
from db import q, ex, now, audit, diff
from security import password_problem
from lifecycle import MGMT_LABEL

EMP_COLS = ["員工編號", "員工姓名", "部門", "地區", "職稱", "到職日", "員工狀態", "初始密碼"]
EMP_REQUIRED = ["員工編號", "員工姓名"]
OFFER_COLS = ["廠商編號", "福利分類", "地區", "優惠名稱", "優惠內容", "使用方式", "注意事項", "開始日期",
              "結束日期", "適用門市", "官方網站", "Google Map", "管理狀態"]
OFFER_REQUIRED = ["廠商編號", "福利分類", "地區", "優惠名稱", "優惠內容", "開始日期", "結束日期"]
STATUS_MAP = {"草稿": "draft", "待審核": "pending"}


class ImportFormatError(Exception):
    pass


def cell_str(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, (datetime, date)):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def parse_date(v):
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = cell_str(v)
    if not s:
        return ""
    m = re.fullmatch(r"(\d{4})[-/.年](\d{1,2})[-/.月](\d{1,2})日?", s)
    if not m:
        return None
    try:
        return date(int(m[1]), int(m[2]), int(m[3])).isoformat()
    except ValueError:
        return None


def read_sheet(fileobj, columns, required):
    """回傳 [(列號, {欄位: 原始值})]。格式不符丟 ImportFormatError。"""
    try:
        wb = openpyxl.load_workbook(fileobj, data_only=True, read_only=True)
    except Exception:
        raise ImportFormatError("檔案無法讀取，請確認是 .xlsx 格式的 Excel 檔案。")
    ws = wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    try:
        header = [cell_str(h) for h in next(rows)]
    except StopIteration:
        raise ImportFormatError("檔案是空白的。")
    missing = [c for c in required if c not in header]
    if missing:
        raise ImportFormatError("缺少必要欄位：" + "、".join(missing) + "。請使用系統提供的範本。")
    idx = {c: header.index(c) for c in columns if c in header}
    out = []
    for n, r in enumerate(rows, start=2):
        if r is None or all(cell_str(x) == "" for x in r):
            continue
        out.append((n, {c: (r[i] if i < len(r) else None) for c, i in idx.items()}))
    if not out:
        raise ImportFormatError("檔案中沒有資料列。")
    if len(out) > 5000:
        raise ImportFormatError("單次最多匯入 5000 筆，請分批上傳。")
    return out


def template_bytes(columns, sample):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(columns)
    ws.append(sample)
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


# ---------------------------------------------------------------- 員工
def validate_employees(rows):
    """rows: [(列號, raw)] → [{row, data, action, errors}]"""
    regions = {r["name"]: r["id"] for r in q("SELECT id,name FROM regions WHERE active=1")}
    existing = {r["emp_no"]: r for r in q("SELECT * FROM employees")}
    seen, out = set(), []
    for n, raw in rows:
        e = []
        d = {c: cell_str(raw.get(c)) for c in EMP_COLS}
        d["到職日"] = parse_date(raw.get("到職日"))
        if not d["員工編號"]:
            e.append("員工編號未填")
        elif not re.fullmatch(r"[A-Za-z0-9_-]{1,20}", d["員工編號"]):
            e.append("員工編號格式不正確（僅限英數字、底線、連字號，20 字內）")
        elif d["員工編號"] in seen:
            e.append("員工編號在檔案中重複")
        seen.add(d["員工編號"])
        if not d["員工姓名"]:
            e.append("員工姓名未填")
        if d["地區"] and d["地區"] not in regions:
            e.append(f"地區「{d['地區']}」不存在")
        if d["到職日"] is None:
            e.append("到職日格式錯誤（請用 YYYY/MM/DD）")
            d["到職日"] = ""
        if d["員工狀態"] and d["員工狀態"] not in ("在職", "留停", "離職"):
            e.append("員工狀態須為 在職／留停／離職")
        cur = existing.get(d["員工編號"])
        action = "update" if cur else "add"
        if action == "add":
            if not d["初始密碼"]:
                e.append("新增員工必須填寫初始密碼")
            elif password_problem(d["初始密碼"]):
                e.append("初始密碼" + password_problem(d["初始密碼"]).replace("密碼", "", 1))
        out.append({"row": n, "data": d, "action": "error" if e else action, "errors": e})
    return out


def apply_employees(items, role_employee_id=1):
    region_ids = {r["name"]: r["id"] for r in q("SELECT id,name FROM regions")}
    t = now()
    added = updated = 0
    for it in items:
        d = it["data"]
        rid = region_ids.get(d["地區"]) if d["地區"] else None
        status = d["員工狀態"] or "在職"
        cur = q("SELECT * FROM employees WHERE emp_no=?", (d["員工編號"],), one=True)
        if cur is None:
            ex("INSERT INTO employees(emp_no,name,dept,title,region_id,hire_date,status,active,role_id,password_hash,"
               "must_change_pw,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,1,?,?)",
               (d["員工編號"], d["員工姓名"], d["部門"], d["職稱"], rid, d["到職日"] or None, status,
                0 if status == "離職" else 1, role_employee_id, generate_password_hash(d["初始密碼"]), t, t))
            added += 1
            audit("員工匯入新增", "employee", d["員工編號"], f'{d["員工姓名"]}（{d["員工編號"]}）')
        else:
            # 不更動角色、不由匯入重設密碼（密碼只能由管理者另行「重設」）
            new = {"name": d["員工姓名"], "dept": d["部門"], "title": d["職稱"], "region_id": rid,
                   "hire_date": d["到職日"] or None, "status": status,
                   "active": 0 if status == "離職" else cur["active"]}
            ch = diff(dict(cur), new)
            if ch:
                ex("UPDATE employees SET name=?,dept=?,title=?,region_id=?,hire_date=?,status=?,active=?,updated_at=? WHERE id=?",
                   (*new.values(), t, cur["id"]))
                audit("員工匯入更新", "employee", cur["id"], f'{d["員工姓名"]}（{d["員工編號"]}）', changes=ch)
            updated += 1
    return added, updated


# ---------------------------------------------------------------- 優惠
def validate_offers(rows):
    cats = {r["name"]: r["id"] for r in q("SELECT id,name FROM categories WHERE active=1")}
    regs = {r["name"]: r["id"] for r in q("SELECT id,name FROM regions WHERE active=1")}
    vendors = {r["code"]: r for r in q("SELECT * FROM vendors")}
    stores = {}
    for s in q("SELECT * FROM stores WHERE status='active'"):
        stores.setdefault(s["vendor_id"], {})[s["name"]] = s["id"]
    out, in_file = [], {}
    for n, raw in rows:
        e = []
        d = {c: cell_str(raw.get(c)) for c in OFFER_COLS}
        for c in OFFER_REQUIRED:
            if c not in ("開始日期", "結束日期") and not d[c]:
                e.append(f"{c}未填")
        ven = vendors.get(d["廠商編號"]) if d["廠商編號"] else None
        if d["廠商編號"] and not ven:
            e.append(f"廠商編號「{d['廠商編號']}」不存在")
        elif ven and ven["status"] != "active":
            e.append(f"廠商「{ven['name']}」已停用，不可建立新優惠")
        if d["福利分類"] and d["福利分類"] not in cats:
            e.append(f"福利分類「{d['福利分類']}」無效")
        reg_names = [x for x in re.split(r"[,，、;；\n]+", d["地區"]) if x.strip()]
        reg_names = [x.strip() for x in reg_names]
        bad = [x for x in reg_names if x not in regs]
        if bad:
            e.append("地區無效：" + "、".join(bad))
        sd, ed = parse_date(raw.get("開始日期")), parse_date(raw.get("結束日期"))
        if sd == "":
            e.append("開始日期未填")
        elif sd is None:
            e.append("開始日期格式錯誤（請用 YYYY/MM/DD）")
        if ed == "":
            e.append("結束日期未填")
        elif ed is None:
            e.append("結束日期格式錯誤（請用 YYYY/MM/DD）")
        if sd and ed and sd > ed:
            e.append("開始日期不得晚於結束日期")
        st = d["管理狀態"] or "草稿"
        if st not in STATUS_MAP:
            e.append("管理狀態僅能匯入「草稿」或「待審核」（發布須經審核流程）")
        store_names = [x.strip() for x in re.split(r"[,，、;；\n]+", d["適用門市"]) if x.strip()]
        if ven and store_names:
            bad = [x for x in store_names if x not in stores.get(ven["id"], {})]
            if bad:
                e.append("門市不存在：" + "、".join(bad))
        action = "add"
        if not e and ven:
            key = (ven["id"], d["優惠名稱"])
            fk = (key, sd)
            if fk in in_file:
                e.append("檔案中有重複的優惠（同廠商、同名稱、同開始日期）")
            else:
                action, msg = _match_existing(ven["id"], d["優惠名稱"], sd, ed, in_file.get(key, []))
                if msg:
                    e.append(msg)
            if not e:
                in_file[fk] = True
                in_file.setdefault(key, []).append((sd, ed))
        d["開始日期"], d["結束日期"] = sd or "", ed or ""
        d["地區"] = "、".join(reg_names)
        out.append({"row": n, "data": d, "action": "error" if e else action, "errors": e})
    return out


def _match_existing(vendor_id, name, sd, ed, pending_periods):
    vs = q("SELECT v.* FROM offer_versions v JOIN offers o ON o.id=v.offer_id "
           "WHERE o.vendor_id=? AND v.name=? AND v.mgmt_status!='archived'", (vendor_id, name))
    for v in vs:
        if v["start_date"] == sd:
            if v["mgmt_status"] in ("draft", "pending"):
                return "update", None
            return "error", f"已有相同期間的版本（{MGMT_LABEL[v['mgmt_status']]}），不可由匯入覆蓋；請使用「續約」建立新版本"
    periods = [(v["start_date"], v["end_date"]) for v in vs] + list(pending_periods)
    for a, b in periods:
        if sd <= b and a <= ed:
            return "error", f"期間與同名優惠的既有版本（{a}～{b}）重疊"
    return "add", None


def apply_offers(items, actor):
    from views_admin_helpers import create_offer_version, next_code
    cats = {r["name"]: r["id"] for r in q("SELECT id,name FROM categories")}
    regs = {r["name"]: r["id"] for r in q("SELECT id,name FROM regions")}
    added = updated = 0
    for it in items:
        d = it["data"]
        ven = q("SELECT * FROM vendors WHERE code=?", (d["廠商編號"],), one=True)
        reg_ids = [regs[x] for x in d["地區"].split("、") if x]
        store_ids = []
        for sn in [x.strip() for x in re.split(r"[,，、;；\n]+", d["適用門市"]) if x.strip()]:
            s = q("SELECT id FROM stores WHERE vendor_id=? AND name=?", (ven["id"], sn), one=True)
            store_ids.append(s["id"])
        # 官方網站／Google Map：僅在廠商資料為空時補入，不覆蓋既有資料
        for col, f in (("官方網站", "website"), ("Google Map", "map_url")):
            if d[col] and d[col].lower().startswith(("http://", "https://")) and not ven[f]:
                ex(f"UPDATE vendors SET {f}=?, updated_at=? WHERE id=?", (d[col], now(), ven["id"]))
        fields = dict(name=d["優惠名稱"], category_id=cats[d["福利分類"]], content=d["優惠內容"],
                      usage=d["使用方式"], notes=d["注意事項"], start_date=d["開始日期"], end_date=d["結束日期"],
                      mgmt_status=STATUS_MAP[d["管理狀態"] or "草稿"])
        existing = q("SELECT v.id FROM offer_versions v JOIN offers o ON o.id=v.offer_id WHERE o.vendor_id=? AND v.name=? "
                     "AND v.start_date=? AND v.mgmt_status IN ('draft','pending')",
                     (ven["id"], d["優惠名稱"], d["開始日期"]), one=True)
        if existing:
            old = dict(q("SELECT * FROM offer_versions WHERE id=?", (existing["id"],), one=True))
            ex("UPDATE offer_versions SET name=?,category_id=?,content=?,usage=?,notes=?,end_date=?,mgmt_status=?,"
               "updated_by=?,updated_at=? WHERE id=?",
               (fields["name"], fields["category_id"], fields["content"], fields["usage"], fields["notes"],
                fields["end_date"], fields["mgmt_status"], actor["id"], now(), existing["id"]))
            ex("DELETE FROM version_regions WHERE version_id=?", (existing["id"],))
            ex("DELETE FROM version_stores WHERE version_id=?", (existing["id"],))
            for r in reg_ids:
                ex("INSERT INTO version_regions VALUES(?,?)", (existing["id"], r))
            for s in store_ids:
                ex("INSERT INTO version_stores VALUES(?,?)", (existing["id"], s))
            audit("優惠匯入更新", "offer_version", existing["id"], d["優惠名稱"], changes=diff(old, fields))
            updated += 1
        else:
            o = q("SELECT o.id FROM offers o JOIN offer_versions v ON v.offer_id=o.id "
                  "WHERE o.vendor_id=? AND v.name=? LIMIT 1", (ven["id"], d["優惠名稱"]), one=True)
            vid = create_offer_version(o["id"] if o else None, ven["id"], fields, reg_ids, store_ids, actor)
            audit("優惠匯入新增", "offer_version", vid, d["優惠名稱"])
            added += 1
    return added, updated


def errors_workbook(errors, columns):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "錯誤清單"
    ws.append(["Excel 列號", "錯誤原因"] + columns)
    for it in errors:
        ws.append([it["row"], "；".join(it["errors"])] + [it["data"].get(c, "") for c in columns])
    bio = io.BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio
